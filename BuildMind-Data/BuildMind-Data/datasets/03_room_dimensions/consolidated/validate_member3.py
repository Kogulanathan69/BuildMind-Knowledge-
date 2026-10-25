#!/usr/bin/env python3
"""Member 3 consolidation validator (BuildMind).

Run from anywhere:
    python validate_member3.py

Needs:
    Python 3.9+
    openpyxl

Exit code:
    0 = every check passed
    1 = at least one FAIL

What it does:
1. Original-file SHA-256 hashes against the manifest recorded before consolidation
2. Every sheet of every original is classified
3. Row-count reconciliation per sheet and in total
4. Required fields, enumerations and policy fields
5. Unit validation
6. Source traceability
7. Cross-version consistency
8. Foreign-reference scan

Important:
This checks internal consistency and provenance only.

It does NOT verify:
- Gazette PDF correctness
- Current legal applicability
- Local authority applicability

Those checks remain PENDING.
"""

import csv
import hashlib
import os
import re
import sys
import collections

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


HERE = os.path.dirname(os.path.abspath(__file__))

ORIG = os.path.normpath(
    os.path.join(HERE, "..", "recovered_versions")
)

TREE = os.path.normpath(
    os.path.join(HERE, "..")
)

RULES_CSV = os.path.join(
    HERE,
    "member3_consolidated_rules.csv"
)

CMP_CSV = os.path.join(
    HERE,
    "member3_cross_version_comparison.csv"
)


# ============================================================
# ORIGINAL FILE HASH MANIFEST
# ============================================================

EXPECTED_SHA256 = {
    "recovered_versions/BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx":
        "4ad3b5789b72c7bf229517cecaf34347466dd0fa302b1e455c71d27f84a176de",

    "recovered_versions/BuildMind_Member3_Gazette_Verified_Dataset.xlsx":
        "6b55339e6a94168d85985229e0cba7e07d32baa5196a056193289920ee19e436",

    "recovered_versions/BuildMind_Member3_Research_Starter.xlsx":
        "39f25a7803404b44a333515e3cfe931d8b7e2933c61b3c9f7f39067fbfa158f0",

    "recovered_versions/BuildMind_Member3_Sri_Lanka_Room_Standards_Extract.xlsx":
        "7e7a6a54690e39093607b6f965081f16a3978245d33423ea0c172a463dc771f2",

    "recovered_versions/BuildMind_Member3_Sri_Lanka_Room_Standards_Stage2.xlsx":
        "e061e77e7512fef4e343291b51e416a9eea0c4c017acf5886071a5c3d975173c",

    "recovered_versions/BuildMind_Member3_Stage3_Concept_Layout_Validation.xlsx":
        "7096e30c635e240b38de260fd391f4ca31ecb918d1e4cb60d71bc33fbfa95de8",

    "recovered_versions/BuildMind_Member3_Stage4_Verified_Research.xlsx":
        "68c6be601a04929975e5e3d76dfbd2d7725173e07b31cd0ec2c7b9afa30e0808",

    "recovered_versions/BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx":
        "da1d25fae9835854bada6b4c4b8739e3779f05009f2169a972fa2b9ca6726a7c",

    "recovered_versions/BuildMind_Member3_Stage6_Gazette_Verification_Audit.xlsx":
        "0b98ebc4fc3a73803619d133a0acb1d29c732e325f4fee6c89d322eee1b37a86",

    "recovered_versions/BuildMind_Member3_Stage9_Gazette_Text_Verified.xlsx":
        "c2d3a078f179272bcacf61196f89b862d47c24c8225277670bfb279401ed0523",

    "recovered_versions/BuildMind_Member3_room_standards.csv":
        "d70cb99d735bd76e8af86d7a06e6633538479a3e243272b353444fa6a92eef3d",

    "recovered_versions/BuildMind_NBRO_Reference_Plan_Extraction_v1.xlsx":
        "8b80bc58a9113e7a5354e5dd12a20077ca4da68cdc51c22370799d62eff17f57",

    "recovered_versions/BuildMind_Stage8_2_Updated_Room_Dimensions.xlsx":
        "d99a0f9e7a379c3e9442bdd3b30d99256e6344f8ae810542b0b390e943084d2c",

    "recovered_versions/BuildMind_Stage8_Sri_Lankan_Floor_Plans_Initial_Extraction.xlsx":
        "1b54bb644da7197105b84d9c227ac509635798d43538b0b046d920a1cf4d66cb",
}


# ============================================================
# RULE SHEETS
# ============================================================

RULES = {
    (
        "BuildMind_Member3_Research_Starter.xlsx",
        "UDA_Minimums",
    ): {
        "kind": "unpivot",
        "id": "record_id",
        "subj": "space_type",
        "metrics": [
            ("minimum_area_m2", "m2"),
            ("minimum_clear_width_m", "m"),
        ],
    },

    (
        "BuildMind_Member3_room_standards.csv",
        "(csv)",
    ): {
        "kind": "unpivot",
        "id": "record_id",
        "subj": "space_type",
        "metrics": [
            ("minimum_area_m2", "m2"),
            ("minimum_clear_width_m", "m"),
        ],
    },

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Extract.xlsx",
        "SL_Verified_Extract",
    ): {
        "kind": "row",
        "id": "Rule_ID",
        "subj": "Space_Type",
        "val": "Value",
        "unit": "Unit",
    },

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Stage2.xlsx",
        "SL_Verified_Extract",
    ): {
        "kind": "row",
        "id": "Rule_ID",
        "subj": "Space_Type",
        "val": "Value",
        "unit": "Unit",
    },

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Stage2.xlsx",
        "Lighting_Ventilation",
    ): {
        "kind": "row",
        "id": "Rule_ID",
        "subj": "Element",
        "val": "Value",
        "unit": "Unit",
    },

    (
        "BuildMind_Member3_Gazette_Verified_Dataset.xlsx",
        "Verified_2021_Rules",
    ): {
        "kind": "row",
        "id": "Rule ID",
        "subj": "Applicable room / category",
        "val": "Threshold",
        "unit": "Unit",
    },

    (
        "BuildMind_Member3_Stage4_Verified_Research.xlsx",
        "Verified_Reference_Rules",
    ): {
        "kind": "row",
        "id": "rule_id",
        "subj": "space_category",
        "val": "value",
        "unit": "unit",
    },

    (
        "BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx",
        "Rule_Verification",
    ): {
        "kind": "row",
        "id": "rule_id",
        "subj": "room_type",
        "val": "claimed_value",
        "unit": "unit",
    },

    (
        "BuildMind_Member3_Stage6_Gazette_Verification_Audit.xlsx",
        "Rule_Audit",
    ): {
        "kind": "row",
        "id": "rule_id",
        "subj": "room_type",
        "val": "candidate_value",
        "unit": "unit",
    },

    (
        "BuildMind_Member3_Stage9_Gazette_Text_Verified.xlsx",
        "Room_Legal_Source",
    ): {
        "kind": "unpivot",
        "id": "Rule_ID",
        "subj": "Room category (exact Gazette wording)",
        "metrics": [
            ("Min area m²", "m2"),
            ("Min width m", "m"),
        ],
    },

    (
        "BuildMind_Member3_Stage9_Gazette_Text_Verified.xlsx",
        "Stairs_Legal_Source",
    ): {
        "kind": "unpivot",
        "id": "Rule_ID",
        "subj": "Gazette category",
        "metrics": [
            ("Min stair width mm", "mm"),
            ("Min headroom m", "m"),
            ("Max riser mm", "mm"),
            ("Min tread mm", "mm"),
        ],
    },

    (
        "BuildMind_Member3_Stage9_Gazette_Text_Verified.xlsx",
        "Room_Heights",
    ): {
        "kind": "unpivot",
        "id": "Rule_ID",
        "subj": "Room categories",
        "metrics": [
            ("Min average height m", "m"),
            ("Min height at lowest point m", "m"),
        ],
    },
}


# ============================================================
# EXCLUDED / NON-RULE SHEETS
# ============================================================

EXCLUDED = {
    (
        "BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx",
        "Source_Register",
    ): "research_literature_non_statutory",

    (
        "BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx",
        "Room_Dimensions",
    ): "research_literature_non_statutory",

    (
        "BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx",
        "Zone_Configurations",
    ): "research_literature_non_statutory",

    (
        "BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx",
        "Design_Findings",
    ): "research_literature_non_statutory",

    (
        "BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx",
        "Legal_Separation",
    ): "research_literature_non_statutory",

    (
        "BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx",
        "Research_Gaps",
    ): "research_literature_non_statutory",

    (
        "BuildMind_Member3_Final_4_Paper_Research_Dataset.xlsx",
        "README",
    ): "research_literature_non_statutory",

    (
        "BuildMind_Member3_Gazette_Verified_Dataset.xlsx",
        "AI_Validation_Examples",
    ): "illustrative_test_cases",

    (
        "BuildMind_Member3_Gazette_Verified_Dataset.xlsx",
        "Read_Me",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Research_Starter.xlsx",
        "Research_Backlog",
    ): "backlog_template",

    (
        "BuildMind_Member3_Research_Starter.xlsx",
        "Sources",
    ): "source_register",

    (
        "BuildMind_Member3_Research_Starter.xlsx",
        "README",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Extract.xlsx",
        "README",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Extract.xlsx",
        "Research_Gaps",
    ): "research_gaps",

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Stage2.xlsx",
        "README",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Stage2.xlsx",
        "Research_Gaps",
    ): "research_gaps",

    (
        "BuildMind_Member3_Sri_Lanka_Room_Standards_Stage2.xlsx",
        "AI_Validation_Examples",
    ): "illustrative_test_cases",

    (
        "BuildMind_Member3_Stage3_Concept_Layout_Validation.xlsx",
        "Concept_Room_Sizes",
    ): "prototype_assumption_concept",

    (
        "BuildMind_Member3_Stage3_Concept_Layout_Validation.xlsx",
        "Furniture_Placeholders",
    ): "prototype_assumption_concept",

    (
        "BuildMind_Member3_Stage3_Concept_Layout_Validation.xlsx",
        "Concept_Calculator",
    ): "prototype_assumption_concept",

    (
        "BuildMind_Member3_Stage3_Concept_Layout_Validation.xlsx",
        "AI_Validation_Rules",
    ): "validation_logic_not_data",

    (
        "BuildMind_Member3_Stage3_Concept_Layout_Validation.xlsx",
        "README",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Stage4_Verified_Research.xlsx",
        "Design_Research_Backlog",
    ): "research_backlog",

    (
        "BuildMind_Member3_Stage4_Verified_Research.xlsx",
        "AI_Validation_Test_Cases",
    ): "illustrative_test_cases",

    (
        "BuildMind_Member3_Stage4_Verified_Research.xlsx",
        "README_Sources",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx",
        "Source_Register",
    ): "source_register",

    (
        "BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx",
        "Jurisdiction_Check",
    ): "process_checklist",

    (
        "BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx",
        "Open_Research_Tasks",
    ): "process_checklist",

    (
        "BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx",
        "AI_Rule_Schema",
    ): "schema_definition",

    (
        "BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx",
        "Validation_Examples",
    ): "illustrative_test_cases",

    (
        "BuildMind_Member3_Stage5_Sri_Lanka_Compliance_Research.xlsx",
        "README",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Stage6_Gazette_Verification_Audit.xlsx",
        "Source_Evidence",
    ): "source_register",

    (
        "BuildMind_Member3_Stage6_Gazette_Verification_Audit.xlsx",
        "Legal_Gates",
    ): "process_checklist",

    (
        "BuildMind_Member3_Stage6_Gazette_Verification_Audit.xlsx",
        "Next_Actions",
    ): "process_checklist",

    (
        "BuildMind_Member3_Stage6_Gazette_Verification_Audit.xlsx",
        "Engine_Decisions",
    ): "illustrative_test_cases",

    (
        "BuildMind_Member3_Stage6_Gazette_Verification_Audit.xlsx",
        "README",
    ): "readme_metadata",

    (
        "BuildMind_Member3_Stage9_Gazette_Text_Verified.xlsx",
        "Research_Gaps",
    ): "research_gaps",

    (
        "BuildMind_Member3_Stage9_Gazette_Text_Verified.xlsx",
        "Source_Register",
    ): "source_register",

    (
        "BuildMind_Member3_Stage9_Gazette_Text_Verified.xlsx",
        "README",
    ): "readme_metadata",

    (
        "BuildMind_NBRO_Reference_Plan_Extraction_v1.xlsx",
        "Drawing_Register",
    ): "plan_observations_examples",

    (
        "BuildMind_NBRO_Reference_Plan_Extraction_v1.xlsx",
        "Room_Extraction",
    ): "plan_observations_examples",

    (
        "BuildMind_NBRO_Reference_Plan_Extraction_v1.xlsx",
        "Opening_Samples",
    ): "plan_observations_examples",

    (
        "BuildMind_NBRO_Reference_Plan_Extraction_v1.xlsx",
        "Data_Rules",
    ): "readme_metadata",

    (
        "BuildMind_Stage8_2_Updated_Room_Dimensions.xlsx",
        "Plan_Register",
    ): "plan_observations_examples",

    (
        "BuildMind_Stage8_2_Updated_Room_Dimensions.xlsx",
        "Room_Extracts",
    ): "plan_observations_examples",

    (
        "BuildMind_Stage8_2_Updated_Room_Dimensions.xlsx",
        "Review_Queue",
    ): "plan_observations_examples",

    (
        "BuildMind_Stage8_2_Updated_Room_Dimensions.xlsx",
        "Data_Rules",
    ): "readme_metadata",

    (
        "BuildMind_Stage8_Sri_Lankan_Floor_Plans_Initial_Extraction.xlsx",
        "Plan_Register",
    ): "plan_observations_examples",

    (
        "BuildMind_Stage8_Sri_Lankan_Floor_Plans_Initial_Extraction.xlsx",
        "Room_Extracts",
    ): "plan_observations_examples",

    (
        "BuildMind_Stage8_Sri_Lankan_Floor_Plans_Initial_Extraction.xlsx",
        "Review_Queue",
    ): "plan_observations_examples",

    (
        "BuildMind_Stage8_Sri_Lankan_Floor_Plans_Initial_Extraction.xlsx",
        "Data_Rules",
    ): "readme_metadata",
}


# ============================================================
# VALIDATION CONSTANTS
# ============================================================

ALLOWED_UNIT_NORM = {
    "m",
    "m2",
    "mm",
    "count",
    "boolean",
    "ratio",
}

ALLOWED_UNIT_SI = {
    "m",
    "m2",
    "count",
    "boolean",
    "ratio",
}

TEXT_ENUM = {
    "TEXT_VERIFIED_PER_ORIGIN",
    "SOURCE_EXTRACTED",
    "PENDING_TEXT_RECHECK",
    "NOT_VERIFIED",
    "NOT_STATED",
}

REQUIRED = [
    "consolidated_id",
    "origin_stage",
    "origin_workbook",
    "origin_sheet",
    "origin_row",
    "origin_rule_id",
    "origin_rule_id_cell",
    "origin_value_cell",
    "origin_value_raw",
    "origin_unit_raw",
    "origin_subject_raw",
    "subject_norm",
    "parameter_norm",
    "value",
    "unit_norm",
    "value_si",
    "unit_si",
    "rule_type",
    "severity",
    "source_document",
    "region",
    "text_verification_status",
    "applicability_status",
    "reviewer_status",
    "independent_pdf_check_by_this_audit",
    "cmp_key",
]

FOREIGN = re.compile(
    r"\b(ICC|IRC|IBC|NBC|Neufert|Eurocode|BS\s?\d|ASHRAE|NFPA|UBC)\b",
    re.I,
)


# ============================================================
# HELPERS
# ============================================================

def param_unit_ok(parameter, unit):
    if parameter in (
        "min_floor_area",
        "airwell_min_area",
        "terrace_house_min_net_vent_area",
    ):
        return unit == "m2"

    if (
        parameter.endswith("_ratio")
        or parameter.endswith("_fraction")
        or parameter in (
            "window_min_opening_area_ratio",
            "window_min_openable_fraction",
            "balcony_min_front_opening_height_fraction",
        )
    ):
        return unit == "ratio"

    if parameter == "dwelling_min_water_closets":
        return unit == "count"

    if parameter == "kitchen_sink_required":
        return unit == "boolean"

    return unit == "m"


results = []
fails = 0


def ok(name, condition, detail=""):
    global fails

    results.append(
        (
            name,
            condition,
            detail,
        )
    )

    print(
        ("PASS  " if condition else "FAIL  ")
        + name
        + (
            ("  -> " + detail)
            if detail
            else ""
        )
    )

    if not condition:
        fails += 1


def S(value):
    if value is None:
        return ""

    return str(value).strip()


# ============================================================
# SHA-256
#
# Important Windows fix:
#
# Git may check text CSV files out using CRLF on Windows.
# The original audited CSV hash was calculated using LF.
#
# For CSV files only, CRLF is normalized to LF before hashing.
#
# XLSX files and other binary files remain byte-for-byte hashed.
# ============================================================

def sha(path):
    with open(path, "rb") as file:
        data = file.read()

    if path.lower().endswith(".csv"):
        data = data.replace(
            b"\r\n",
            b"\n",
        )

    return hashlib.sha256(
        data
    ).hexdigest()


# ============================================================
# LOAD ORIGINAL DATA
# ============================================================

_wb = {}


def book(filename):
    if filename not in _wb:
        _wb[filename] = load_workbook(
            os.path.join(
                ORIG,
                filename,
            ),
            data_only=True,
        )

    return _wb[filename]


def csv_lines():
    output = {}

    path = os.path.join(
        ORIG,
        "BuildMind_Member3_room_standards.csv",
    )

    with open(
        path,
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.reader(
            file
        )

        header = next(
            reader
        )

        for row in reader:
            if any(
                value.strip()
                for value in row
            ):
                output[
                    reader.line_num
                ] = row

    return header, output


def sheet_data(filename, sheet):
    """Return header list and {row_number: values}."""

    if filename.endswith(
        ".csv"
    ):
        header, data = csv_lines()

        return (
            header,
            data,
        )

    worksheet = book(
        filename
    )[sheet]

    header = None
    data = {}

    for row_number in range(
        1,
        worksheet.max_row + 1,
    ):
        values = [
            worksheet.cell(
                row_number,
                column,
            ).value
            for column in range(
                1,
                worksheet.max_column + 1,
            )
        ]

        if not any(
            value is not None
            and str(
                value
            ).strip()
            for value in values
        ):
            continue

        if header is None:
            header = [
                S(
                    value
                )
                for value in values
            ]

            continue

        data[
            row_number
        ] = values

    return (
        header,
        data,
    )


def cell_text(
    filename,
    sheet,
    reference,
):
    if filename.endswith(
        ".csv"
    ):
        _, data = csv_lines()

        column = re.match(
            r"[A-Z]+",
            reference,
        ).group()

        row = int(
            reference[
                len(
                    column
                ):
            ]
        )

        from openpyxl.utils import column_index_from_string

        index = (
            column_index_from_string(
                column
            )
            - 1
        )

        if (
            row in data
            and index
            < len(
                data[
                    row
                ]
            )
        ):
            return S(
                data[
                    row
                ][
                    index
                ]
            )

        return ""

    return S(
        book(
            filename
        )[
            sheet
        ][
            reference
        ].value
    )


# ============================================================
# START VALIDATION
# ============================================================

print(
    "Member 3 validation  (originals: %s)\n"
    % ORIG
)


# ============================================================
# CHECK 1
# ORIGINAL FILE HASHES
# ============================================================

tree = {}

for root, _, files in os.walk(
    TREE
):
    for filename in files:

        path = os.path.join(
            root,
            filename,
        )

        relative = os.path.relpath(
            path,
            TREE,
        ).replace(
            os.sep,
            "/",
        )

        if relative.startswith(
            "consolidated/"
        ):
            continue

        tree[
            relative
        ] = path


missing = [
    key
    for key in EXPECTED_SHA256
    if key not in tree
]

extra = [
    key
    for key in tree
    if key not in EXPECTED_SHA256
]

bad = [
    key
    for key in EXPECTED_SHA256
    if (
        key in tree
        and sha(
            tree[
                key
            ]
        )
        != EXPECTED_SHA256[
            key
        ]
    )
]


ok(
    "1 original files present (%d expected)"
    % len(
        EXPECTED_SHA256
    ),
    not missing,
    (
        "missing: %s"
        % missing
        if missing
        else ""
    ),
)

ok(
    "1 no unexpected files beside originals (outside consolidated/)",
    not extra,
    (
        "extra: %s"
        % extra
        if extra
        else ""
    ),
)

ok(
    "1 SHA-256 of every original equals the pre-consolidation manifest",
    not bad,
    (
        "changed: %s"
        % bad
        if bad
        else "%d/%d identical"
        % (
            len(
                EXPECTED_SHA256
            ),
            len(
                EXPECTED_SHA256
            ),
        )
    ),
)


# ============================================================
# CHECK 2
# SHEET CLASSIFICATION
# ============================================================

unclassified = []

classified = (
    set(
        RULES
    )
    |
    set(
        EXCLUDED
    )
)


for filename in sorted(
    os.listdir(
        ORIG
    )
):
    if filename.endswith(
        ".csv"
    ):
        sheets = [
            "(csv)"
        ]

    else:
        sheets = book(
            filename
        ).sheetnames

    for sheet in sheets:
        if (
            filename,
            sheet,
        ) not in classified:
            unclassified.append(
                (
                    filename,
                    sheet,
                )
            )


absent = []

for key in classified:

    filename, sheet = key

    full_path = os.path.join(
        ORIG,
        filename,
    )

    if not os.path.exists(
        full_path
    ):
        absent.append(
            key
        )

        continue

    if (
        not filename.endswith(
            ".csv"
        )
        and sheet
        not in book(
            filename
        ).sheetnames
    ):
        absent.append(
            key
        )


ok(
    "2 every sheet of every original is classified (%d sheets)"
    % len(
        classified
    ),
    not unclassified
    and not absent,
    (
        str(
            unclassified
            + absent
        )
        if (
            unclassified
            or absent
        )
        else ""
    ),
)


# ============================================================
# LOAD CONSOLIDATED OUTPUTS
# ============================================================

with open(
    RULES_CSV,
    encoding="utf-8-sig",
    newline="",
) as file:
    rules = list(
        csv.DictReader(
            file
        )
    )


with open(
    CMP_CSV,
    encoding="utf-8-sig",
    newline="",
) as file:
    cmpr = list(
        csv.DictReader(
            file
        )
    )


# ============================================================
# CHECK 3
# ROW RECONCILIATION
# ============================================================

by_origin = collections.Counter(
    (
        row[
            "origin_workbook"
        ],
        row[
            "origin_sheet"
        ],
    )
    for row in rules
)

expected = {}

excluded_rows = {}


for (
    filename,
    sheet,
), specification in RULES.items():

    header, data = sheet_data(
        filename,
        sheet,
    )

    if specification[
        "kind"
    ] == "row":

        value_index = header.index(
            specification[
                "val"
            ]
        )

        expected[
            (
                filename,
                sheet,
            )
        ] = sum(
            1
            for _,
            values in data.items()
            if S(
                values[
                    value_index
                ]
            )
            != ""
        )

    else:

        indexes = [
            header.index(
                metric
            )
            for metric, _
            in specification[
                "metrics"
            ]
        ]

        expected[
            (
                filename,
                sheet,
            )
        ] = sum(
            1
            for _,
            values in data.items()
            for index in indexes
            if (
                index
                < len(
                    values
                )
                and S(
                    values[
                        index
                    ]
                )
                != ""
            )
        )


bad_counts = [
    (
        key,
        expected[
            key
        ],
        by_origin.get(
            key,
            0,
        ),
    )
    for key in expected
    if expected[
        key
    ]
    != by_origin.get(
        key,
        0,
    )
]


ok(
    "3 per-sheet reconciliation: origin value count == consolidated rows (%d rule sheets)"
    % len(
        expected
    ),
    not bad_counts,
    (
        str(
            bad_counts
        )
        if bad_counts
        else ""
    ),
)


ok(
    "3 total reconciliation: %d expected == %d consolidated"
    % (
        sum(
            expected.values()
        ),
        len(
            rules
        ),
    ),
    (
        sum(
            expected.values()
        )
        ==
        len(
            rules
        )
        ==
        sum(
            by_origin.values()
        )
    ),
)


stray = [
    key
    for key in by_origin
    if key not in RULES
]


ok(
    "3 no consolidated row originates from an excluded/unknown sheet",
    not stray,
    (
        str(
            stray
        )
        if stray
        else ""
    ),
)


total_excluded = 0

for (
    filename,
    sheet,
) in EXCLUDED:

    header, data = sheet_data(
        filename,
        sheet,
    )

    excluded_rows[
        (
            filename,
            sheet,
        )
    ] = len(
        data
    )

    total_excluded += len(
        data
    )


print(
    "      (info) %d sheets / %d data rows are excluded as non-rule data; see audit report"
    % (
        len(
            EXCLUDED
        ),
        total_excluded,
    )
)


# ============================================================
# CHECK 4
# REQUIRED FIELDS / POLICY
# ============================================================

missing_required = [
    (
        row[
            "consolidated_id"
        ],
        column,
    )
    for row in rules
    for column in REQUIRED
    if S(
        row.get(
            column
        )
    )
    == ""
]


ok(
    "4 required fields non-empty on all rows",
    not missing_required,
    (
        str(
            missing_required[
                :5
            ]
        )
        if missing_required
        else ""
    ),
)


expected_ids = [
    "M3-%04d"
    % index
    for index in range(
        1,
        len(
            rules
        )
        + 1,
    )
]


ok(
    "4 consolidated_id unique and sequential",
    [
        row[
            "consolidated_id"
        ]
        for row in rules
    ]
    ==
    expected_ids,
)


ok(
    "4 applicability_status is PENDING on every row (text verified != currently applicable)",
    all(
        row[
            "applicability_status"
        ]
        ==
        "PENDING"
        for row in rules
    ),
)


ok(
    "4 text_verification_status within allowed set",
    all(
        row[
            "text_verification_status"
        ]
        in TEXT_ENUM
        for row in rules
    ),
)


ok(
    "4 reviewer_status NOT_RECORDED on every row (no origin file records a reviewer)",
    all(
        row[
            "reviewer_status"
        ]
        ==
        "NOT_RECORDED"
        for row in rules
    ),
)


ok(
    "4 severity NOT_ASSIGNED on every row (origin does not assign hard/soft)",
    all(
        row[
            "severity"
        ]
        ==
        "NOT_ASSIGNED"
        for row in rules
    ),
)


ok(
    "4 independent PDF check is NOT_PERFORMED on every row",
    all(
        row[
            "independent_pdf_check_by_this_audit"
        ].startswith(
            "NOT_PERFORMED"
        )
        for row in rules
    ),
)


flag_mismatch = [
    row[
        "consolidated_id"
    ]
    for row in rules
    if (
        (
            S(
                row[
                    "origin_clause_raw"
                ]
            )
            == ""
        )
        != (
            "MISSING_CLAUSE"
            in row[
                "flags"
            ].split(
                ";"
            )
        )
        or
        (
            (
                S(
                    row[
                        "origin_page_raw"
                    ]
                )
                == ""
            )
            != (
                "MISSING_PAGE"
                in row[
                    "flags"
                ].split(
                    ";"
                )
            )
        )
    )
]


ok(
    "4 missing clause/page are explicitly flagged (no silent blanks)",
    not flag_mismatch,
    (
        str(
            flag_mismatch[
                :5
            ]
        )
        if flag_mismatch
        else ""
    ),
)


no_operator = [
    row[
        "consolidated_id"
    ]
    for row in rules
    if (
        row[
            "operator"
        ]
        == ""
        and "not stated"
        not in row[
            "operator_basis"
        ]
    )
]


ok(
    "4 operator blank only where origin gives none and metric is a flag/threshold",
    not no_operator,
)


# ============================================================
# CHECK 5
# UNIT VALIDATION
# ============================================================

bad_units = [
    row[
        "consolidated_id"
    ]
    for row in rules
    if (
        row[
            "unit_norm"
        ]
        not in ALLOWED_UNIT_NORM
        or row[
            "unit_si"
        ]
        not in ALLOWED_UNIT_SI
    )
]


ok(
    "5 units belong to the allowed set",
    not bad_units,
    (
        str(
            bad_units[
                :5
            ]
        )
        if bad_units
        else ""
    ),
)


bad_parameters = [
    (
        row[
            "consolidated_id"
        ],
        row[
            "parameter_norm"
        ],
        row[
            "unit_si"
        ],
    )
    for row in rules
    if not param_unit_ok(
        row[
            "parameter_norm"
        ],
        row[
            "unit_si"
        ],
    )
]


ok(
    "5 parameter/unit compatibility (area=m2, lengths=m, ratios=ratio, ...)",
    not bad_parameters,
    (
        str(
            bad_parameters[
                :5
            ]
        )
        if bad_parameters
        else ""
    ),
)


bad_conversion = [
    row[
        "consolidated_id"
    ]
    for row in rules
    if abs(
        float(
            row[
                "value"
            ]
        )
        *
        (
            0.001
            if row[
                "unit_norm"
            ]
            ==
            "mm"
            else 1
        )
        -
        float(
            row[
                "value_si"
            ]
        )
    )
    >
    1e-9
]


ok(
    "5 value_si recomputed from value and unit (mm -> m x0.001)",
    not bad_conversion,
    (
        str(
            bad_conversion[
                :5
            ]
        )
        if bad_conversion
        else ""
    ),
)


millimetre_rows = sum(
    1
    for row in rules
    if row[
        "unit_norm"
    ]
    ==
    "mm"
)


print(
    "      (info) %d rows were stated in mm and are compared in metres"
    % millimetre_rows
)


bad_operators = [
    row[
        "consolidated_id"
    ]
    for row in rules
    if (
        (
            "_max_"
            in row[
                "parameter_norm"
            ]
            or row[
                "parameter_norm"
            ].startswith(
                "max"
            )
            or row[
                "parameter_norm"
            ]
            ==
            "stair_max_riser"
        )
        and row[
            "operator"
        ]
        !=
        "<="
    )
]


ok(
    "5 maximum-type parameters use operator <= (no min/max inversion)",
    not bad_operators,
    (
        str(
            bad_operators
        )
        if bad_operators
        else ""
    ),
)


# ============================================================
# CHECK 6
# SOURCE TRACEABILITY
# ============================================================

traceability_bad = []


for row_data in rules:

    filename = row_data[
        "origin_workbook"
    ]

    sheet = row_data[
        "origin_sheet"
    ]

    specification = RULES.get(
        (
            filename,
            sheet,
        )
    )

    if not specification:
        traceability_bad.append(
            (
                row_data[
                    "consolidated_id"
                ],
                "unknown origin sheet",
            )
        )

        continue


    origin_row = int(
        row_data[
            "origin_row"
        ]
    )

    value_cell = row_data[
        "origin_value_cell"
    ]

    id_cell = row_data[
        "origin_rule_id_cell"
    ]


    if (
        int(
            re.search(
                r"\d+",
                value_cell,
            ).group()
        )
        !=
        origin_row
        or
        int(
            re.search(
                r"\d+",
                id_cell,
            ).group()
        )
        !=
        origin_row
    ):
        traceability_bad.append(
            (
                row_data[
                    "consolidated_id"
                ],
                "cell row != origin_row",
            )
        )

        continue


    if (
        cell_text(
            filename,
            sheet,
            id_cell,
        )
        !=
        row_data[
            "origin_rule_id"
        ]
    ):
        traceability_bad.append(
            (
                row_data[
                    "consolidated_id"
                ],
                "rule id mismatch",
            )
        )

        continue


    if (
        cell_text(
            filename,
            sheet,
            value_cell,
        )
        !=
        row_data[
            "origin_value_raw"
        ]
    ):
        traceability_bad.append(
            (
                row_data[
                    "consolidated_id"
                ],
                "value mismatch",
            )
        )

        continue


    try:
        if (
            float(
                row_data[
                    "origin_value_raw"
                ]
            )
            !=
            float(
                row_data[
                    "value"
                ]
            )
        ):
            traceability_bad.append(
                (
                    row_data[
                        "consolidated_id"
                    ],
                    "numeric mismatch",
                )
            )

            continue

    except ValueError:
        traceability_bad.append(
            (
                row_data[
                    "consolidated_id"
                ],
                "non-numeric",
            )
        )

        continue


    header, origin_data = sheet_data(
        filename,
        sheet,
    )


    if origin_row not in origin_data:
        traceability_bad.append(
            (
                row_data[
                    "consolidated_id"
                ],
                "row not in origin",
            )
        )

        continue


    subject_index = header.index(
        specification[
            "subj"
        ]
    )


    if (
        S(
            origin_data[
                origin_row
            ][
                subject_index
            ]
        )
        !=
        row_data[
            "origin_subject_raw"
        ]
    ):
        traceability_bad.append(
            (
                row_data[
                    "consolidated_id"
                ],
                "subject mismatch",
            )
        )

        continue


    if specification[
        "kind"
    ] == "row":

        unit_index = header.index(
            specification[
                "unit"
            ]
        )

        if (
            S(
                origin_data[
                    origin_row
                ][
                    unit_index
                ]
            )
            !=
            row_data[
                "origin_unit_raw"
            ]
        ):
            traceability_bad.append(
                (
                    row_data[
                        "consolidated_id"
                    ],
                    "unit mismatch",
                )
            )

            continue


        expected_column = (
            get_column_letter(
                header.index(
                    specification[
                        "val"
                    ]
                )
                + 1
            )
            + str(
                origin_row
            )
        )


        if expected_column != value_cell:
            traceability_bad.append(
                (
                    row_data[
                        "consolidated_id"
                    ],
                    "value column mismatch",
                )
            )

            continue

    else:

        column = re.match(
            r"[A-Z]+",
            value_cell,
        ).group()

        metric_names = [
            metric
            for metric, _
            in specification[
                "metrics"
            ]
        ]

        metric_units = dict(
            specification[
                "metrics"
            ]
        )

        hit = [
            metric
            for metric
            in metric_names
            if get_column_letter(
                header.index(
                    metric
                )
                + 1
            )
            ==
            column
        ]


        if (
            len(
                hit
            )
            !=
            1
            or
            metric_units[
                hit[
                    0
                ]
            ]
            !=
            row_data[
                "origin_unit_raw"
            ]
            or
            hit[
                0
            ]
            !=
            row_data[
                "origin_metric_raw"
            ]
        ):
            traceability_bad.append(
                (
                    row_data[
                        "consolidated_id"
                    ],
                    "metric/unit mismatch",
                )
            )

            continue


ok(
    "6 source traceability: every row re-matches its origin cell (id, value, subject, unit, metric)",
    not traceability_bad,
    (
        str(
            traceability_bad[
                :5
            ]
        )
        if traceability_bad
        else "%d/%d rows verified against originals"
        % (
            len(
                rules
            ),
            len(
                rules
            ),
        )
    ),
)


# ============================================================
# CHECK 7
# CROSS-VERSION CONSISTENCY
# ============================================================

groups = collections.defaultdict(
    list
)


for row in rules:
    groups[
        row[
            "cmp_key"
        ]
    ].append(
        row
    )


recomputed = {}


for key, rows in groups.items():

    stages = {
        row[
            "origin_stage"
        ]
        for row in rows
    }

    values = {
        round(
            float(
                row[
                    "value_si"
                ]
            ),
            9,
        )
        for row in rows
    }


    if len(
        stages
    ) == 1:
        status = "SINGLE_ORIGIN"

    elif len(
        values
    ) == 1:
        status = "CONSISTENT"

    else:
        status = "CONFLICT"


    recomputed[
        key
    ] = (
        status,
        len(
            rows
        ),
        len(
            stages
        ),
    )


ok(
    "7 comparison CSV has one row per comparison key, no duplicates",
    (
        len(
            cmpr
        )
        ==
        len(
            {
                row[
                    "cmp_key"
                ]
                for row in cmpr
            }
        )
        ==
        len(
            groups
        )
    ),
)


comparison_mismatch = [
    row[
        "cmp_key"
    ]
    for row in cmpr
    if recomputed.get(
        row[
            "cmp_key"
        ]
    )
    !=
    (
        row[
            "value_consistency"
        ],
        int(
            row[
                "n_rows"
            ]
        ),
        int(
            row[
                "n_origin_versions"
            ]
        ),
    )
]


ok(
    "7 comparison CSV statuses/counts equal an independent recomputation",
    not comparison_mismatch,
    (
        str(
            comparison_mismatch[
                :5
            ]
        )
        if comparison_mismatch
        else ""
    ),
)


ok(
    "7 comparison rows cover all consolidated rows",
    sum(
        int(
            row[
                "n_rows"
            ]
        )
        for row in cmpr
    )
    ==
    len(
        rules
    ),
)


conflicts = [
    key
    for key, value
    in recomputed.items()
    if value[
        0
    ]
    ==
    "CONFLICT"
]


print(
    "      (finding) value conflicts across versions: %d %s"
    % (
        len(
            conflicts
        ),
        conflicts,
    )
)


print(
    "      (finding) comparison keys: %d consistent, %d single-origin, %d conflict"
    %
    tuple(
        sum(
            1
            for value
            in recomputed.values()
            if value[
                0
            ]
            ==
            status
        )
        for status in (
            "CONSISTENT",
            "SINGLE_ORIGIN",
            "CONFLICT",
        )
    )
)


# ============================================================
# CHECK 8
# FOREIGN REFERENCE SCAN
# ============================================================

foreign_hits = [
    (
        row[
            "consolidated_id"
        ],
        column,
    )
    for row in rules
    for column in (
        "source_document",
        "source_url",
        "origin_clause_raw",
        "source_file_name",
    )
    if FOREIGN.search(
        row.get(
            column,
            "",
        )
        or ""
    )
]


ok(
    "8 no foreign-code/standard reference in source fields",
    not foreign_hits,
    (
        str(
            foreign_hits[
                :5
            ]
        )
        if foreign_hits
        else ""
    ),
)


foreign_note_hits = [
    row[
        "consolidated_id"
    ]
    for row in rules
    if FOREIGN.search(
        " ".join(
            [
                row[
                    "origin_notes_raw"
                ],
                row[
                    "origin_condition_raw"
                ],
                row[
                    "origin_applicability_raw"
                ],
            ]
        )
    )
]


print(
    "      (info) foreign-reference words in origin note fields on %d rows %s"
    % (
        len(
            foreign_note_hits
        ),
        foreign_note_hits[
            :5
        ],
    )
)


# ============================================================
# FINAL SUMMARY
# ============================================================

print(
    "\nSUMMARY: %d checks, %d passed, %d failed"
    % (
        len(
            results
        ),
        len(
            results
        )
        - fails,
        fails,
    )
)


sys.exit(
    1
    if fails
    else 0
)